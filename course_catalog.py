# -*- coding: utf-8 -*-
"""
课程体系展示模块 —— 数据层
==========================

职责：
1. 定义「专业 Major → 课程类别 CourseCategory → 课程 Course」三级数据模型；
2. 解析专业培养方案 Excel（首版：软件工程专业），并通过 Streamlit 缓存机制
   避免重复读盘；
3. 管理多专业注册表（MAJORS），后续新增专业只需在表中登记一条记录；
4. 提供「AI 学习建议」能力：优先调用 DeepSeek 生成，结果持久化到 JSON 缓存，
   未配置 API Key 或调用失败时降级为基于表格字段的规则化建议。

数据来源约定：
    Excel 主表首行表头，之后每行一门课程；「课程类别」为合并单元格（仅每类
    第一行有值），解析时向下填充。
"""

from __future__ import annotations

# 在导入 openpyxl 前关闭字节码落盘，避免受限环境下写 __pycache__ 触发权限错误
import sys

sys.dont_write_bytecode = True

import json
import os
import time
from dataclasses import dataclass, field
from typing import Dict, List, Optional

import streamlit as st

try:
    import openpyxl
except ImportError:  # pragma: no cover - 依赖缺失时给出可操作的提示
    openpyxl = None


# --------------------------------------------------------------------------- #
# 路径与常量
# --------------------------------------------------------------------------- #
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ADVICE_CACHE_FILE = os.path.join(BASE_DIR, "course_advice_cache.json")

# 类别顺序与 Excel 中出现的顺序保持一致（培养路径顺序）
CATEGORY_ORDER = [
    "公共基础课",
    "专业基础课",
    "专业核心课",
    "专业拓展选修课",
    "实践教学环节",
]

# 类别元数据：编码（用于生成课程编号）、图标、教育感浅色配色（取自表格使用说明图例）
CATEGORY_META: Dict[str, Dict[str, str]] = {
    "公共基础课": {
        "code": "GE",
        "icon": "📐",
        # 浅蓝
        "tint": "#EAF2FE", "tint_strong": "#D6E6FD",
        "accent": "#2563EB", "accent_soft": "rgba(37,99,235,.10)",
        "desc": "夯实数理与通识基础，建立理工科思维",
    },
    "专业基础课": {
        "code": "FB",
        "icon": "🧱",
        # 浅绿
        "tint": "#E8F7EE", "tint_strong": "#D2F0DF",
        "accent": "#059669", "accent_soft": "rgba(5,150,105,.10)",
        "desc": "计算机学科核心理论底座，承上启下",
    },
    "专业核心课": {
        "code": "PC",
        "icon": "🎯",
        # 浅橙
        "tint": "#FDF1E3", "tint_strong": "#FAE2C4",
        "accent": "#EA7A0C", "accent_soft": "rgba(234,122,12,.10)",
        "desc": "软件工程方法论与工程能力主干课程",
    },
    "专业拓展选修课": {
        "code": "EE",
        "icon": "🧭",
        # 浅紫
        "tint": "#F1ECFD", "tint_strong": "#E4DAFB",
        "accent": "#7C3AED", "accent_soft": "rgba(124,58,237,.10)",
        "desc": "面向前沿方向与职业兴趣的拓展选修",
    },
    "实践教学环节": {
        "code": "PR",
        "icon": "🛠️",
        # 浅红
        "tint": "#FDECEC", "tint_strong": "#FBDADA",
        "accent": "#E11D48", "accent_soft": "rgba(225,29,72,.10)",
        "desc": "综合运用知识解决真实工程问题",
    },
}


# --------------------------------------------------------------------------- #
# 数据模型
# --------------------------------------------------------------------------- #
@dataclass
class Course:
    """单门课程的完整信息（对应 Excel 一行 17 列）。"""

    seq: int
    category: str
    name: str
    name_en: str = ""
    nature: str = ""
    credits: str = ""
    hours: str = ""
    semester: str = ""
    intro: str = ""
    core_content: str = ""
    objectives: str = ""
    prerequisites: str = ""
    assessment: str = ""
    difficulty: str = ""
    practice_ratio: str = ""
    career: str = ""
    remark: str = ""
    code: str = ""  # 展示用课程编号（按专业+类别+序号生成）

    @property
    def difficulty_level(self) -> int:
        """星级文本（如 ★★★★☆）转数字，用于排序与视觉渲染。"""
        return self.difficulty.count("★")

    @property
    def content_points(self) -> List[str]:
        """核心学习内容按分隔符拆成知识点标签。"""
        raw = self.core_content.replace("，", "、").replace(",", "、").replace("；", "、")
        return [p.strip() for p in raw.split("、") if p.strip()]

    @property
    def prerequisite_list(self) -> List[str]:
        raw = self.prerequisites.replace("，", "、").replace("；", "、").replace(";", "、")
        return [p.strip() for p in raw.split("、") if p.strip() and p.strip() not in ("无", "-")]


@dataclass
class Major:
    """一个专业（培养方案）：对应一个 Excel 工作簿。"""

    id: str
    name: str
    name_en: str
    icon: str
    excel_name: str
    sheet_name: str
    tagline: str
    available: bool = True
    courses: List[Course] = field(default_factory=list)

    @property
    def excel_path(self) -> str:
        return os.path.join(BASE_DIR, self.excel_name)

    @property
    def total_credits(self) -> float:
        total = 0.0
        for c in self.courses:
            try:
                total += float(str(c.credits).strip())
            except (TypeError, ValueError):
                continue
        return total

    def courses_by_category(self, category: str) -> List[Course]:
        return [c for c in self.courses if c.category == category]


# --------------------------------------------------------------------------- #
# 多专业注册表：后续新增专业只需在此登记（并把对应 Excel 放入项目目录）
# --------------------------------------------------------------------------- #
MAJORS: List[Major] = [
    Major(
        id="SE",
        name="软件工程",
        name_en="Software Engineering",
        icon="💻",
        excel_name="软件工程专业课程详细介绍表.xlsx",
        sheet_name="软件工程专业课程详解",
        tagline="系统化呈现软件开发全生命周期的知识体系与实践能力培养路径",
        available=True,
    ),
    # —— 扩展示例（放开后放入对应 Excel 即可上线）——
    # Major(
    #     id="CS", name="计算机科学与技术", name_en="Computer Science",
    #     icon="🖥️", excel_name="计算机科学与技术专业课程详细介绍表.xlsx",
    #     sheet_name="计算机科学与技术专业课程详解",
    #     tagline="聚焦计算理论、系统与算法的宽口径培养", available=True,
    # ),
    Major(
        id="COMING",
        name="更多专业",
        name_en="More Majors",
        icon="✨",
        excel_name="",
        sheet_name="",
        tagline="人工智能、数据科学、网络工程等专业课程体系正在整理中，敬请期待",
        available=False,
    ),
]


def get_major(major_id: str) -> Optional[Major]:
    return next((m for m in MAJORS if m.id == major_id), None)


# --------------------------------------------------------------------------- #
# Excel 解析（带 Streamlit 数据缓存，依赖文件修改时间自动失效）
# --------------------------------------------------------------------------- #
_EXCEL_COLUMNS = {
    "序号": "seq",
    "课程类别": "category",
    "课程名称": "name",
    "英文名称": "name_en",
    "课程性质": "nature",
    "建议学分": "credits",
    "建议学时": "hours",
    "建议开课学期": "semester",
    "课程简介": "intro",
    "核心学习内容": "core_content",
    "能力培养目标": "objectives",
    "建议先修课程": "prerequisites",
    "考核方式": "assessment",
    "难度评级": "difficulty",
    "实践占比": "practice_ratio",
    "对应职业方向": "career",
    "备注": "remark",
}


@st.cache_data(show_spinner=False, max_entries=8)
def _parse_excel(_path: str, sheet_name: str, mtime: float) -> List[dict]:
    """
    读取 Excel 并转成课程字典列表。

    :param mtime: 文件修改时间，作为缓存失效因子（内容更新后自动重新解析）
    """
    if openpyxl is None:
        raise RuntimeError("缺少 openpyxl 依赖，请先执行：pip install openpyxl")
    wb = openpyxl.load_workbook(_path, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    index = {h: i for i, h in enumerate(headers) if h in _EXCEL_COLUMNS}

    result: List[dict] = []
    last_category: Optional[str] = None
    for raw in rows[1:]:
        if not any(v is not None and str(v).strip() for v in raw):
            continue  # 跳过整行空行

        def cell(col_name: str) -> str:
            i = index.get(col_name)
            if i is None or i >= len(raw) or raw[i] is None:
                return ""
            return str(raw[i]).strip()

        category = cell("课程类别") or last_category
        last_category = category
        name = cell("课程名称")
        if not name:
            continue

        item = {field_key: cell(col) for col, field_key in _EXCEL_COLUMNS.items()}
        item["category"] = category
        # 序号转 int（失败时按行号兜底）
        try:
            item["seq"] = int(float(cell("序号")))
        except (TypeError, ValueError):
            item["seq"] = len(result) + 1
        result.append(item)
    return result


def load_major_courses(major: Major) -> Major:
    """
    加载某专业的全部课程（含缓存）。已加载则直接返回，避免重复解析。
    """
    if major.courses or not major.available:
        return major

    try:
        mtime = os.path.getmtime(major.excel_path)
    except OSError:
        mtime = 0.0
    records = _parse_excel(major.excel_path, major.sheet_name, mtime)

    # 类别内序号（用于课程编号的尾号）
    cat_counter: Dict[str, int] = {}
    courses: List[Course] = []
    for rec in records:
        category = rec.get("category") or CATEGORY_ORDER[0]
        cat_counter[category] = cat_counter.get(category, 0) + 1
        meta = CATEGORY_META.get(category)
        code_prefix = meta["code"] if meta else "XX"
        rec["code"] = f"{major.id}-{code_prefix}-{cat_counter[category]:02d}"
        courses.append(Course(**{k: v for k, v in rec.items() if k in Course.__dataclass_fields__}))

    # 按类别顺序、类别内序号稳定排序
    order_map = {c: i for i, c in enumerate(CATEGORY_ORDER)}
    courses.sort(key=lambda c: (order_map.get(c.category, 99), c.seq))
    major.courses = courses
    return major


def find_successors(course: Course, all_courses: List[Course]) -> List[str]:
    """在全部课程中反查：哪些课程把本课程列为先修（用于学习路径串联）。"""
    names = []
    keywords = [course.name]
    # 去掉括号补充，如「面向对象程序设计（Java/Python）」->「面向对象程序设计」
    if "（" in course.name:
        keywords.append(course.name.split("（")[0])
    for other in all_courses:
        if other.name == course.name:
            continue
        prereq = other.prerequisites
        if prereq and any(k and k in prereq for k in keywords):
            names.append(other.name)
    return names


# --------------------------------------------------------------------------- #
# AI 学习建议：DeepSeek 生成 + JSON 磁盘缓存 + 规则化降级
# --------------------------------------------------------------------------- #
ADVICE_SYSTEM_PROMPT = (
    "你是大学软件工程专业的资深学业规划导师，擅长为不同基础的学生制定务实的课程学习方案。"
    "请依据给出的课程信息输出 Markdown，严格包含以下三个小节（标题保持一致）：\n"
    "### 🎯 重要性分析\n说明该课程在整个专业培养体系中的定位、与其他课程的关联、对职业方向的价值；3 条以内。\n"
    "### 🗺️ 学习路径\n按「先修准备 → 本课程重点 → 后续延伸」给出递进路径，步骤清晰；可结合给出的先修课与后续课程。\n"
    "### 💡 学习建议\n给出可执行的学习方法、实践安排与避坑提示，结合课程难度与实践占比；3 条以内。\n"
    "要求：全部简体中文，语言精炼专业，不要空泛套话，不要编造课程信息中没有的事实。"
)


def _advice_cache_key(major_id: str, course_code: str) -> str:
    return f"{major_id}:{course_code}"


def load_advice_cache() -> Dict[str, dict]:
    try:
        with open(ADVICE_CACHE_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (FileNotFoundError, ValueError, OSError):
        return {}


def save_advice_cache(cache: Dict[str, dict]) -> None:
    """原子写入建议缓存。"""
    tmp = ADVICE_CACHE_FILE + ".tmp"
    try:
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False, indent=2)
        os.replace(tmp, ADVICE_CACHE_FILE)
    except OSError:
        pass


def get_cached_advice(major_id: str, course_code: str) -> Optional[str]:
    entry = load_advice_cache().get(_advice_cache_key(major_id, course_code))
    return entry.get("text") if isinstance(entry, dict) else None


def _build_course_context(course: Course, successors: List[str]) -> str:
    lines = [
        f"课程名称：{course.name}（{course.name_en}）",
        f"课程类别：{course.category}　课程性质：{course.nature}",
        f"学分/学时：{course.credits} 学分 / {course.hours} 学时　建议学期：{course.semester}",
        f"课程简介：{course.intro}",
        f"核心学习内容：{course.core_content}",
        f"能力培养目标：{course.objectives}",
        f"建议先修课程：{course.prerequisites}",
        f"考核方式：{course.assessment}　难度评级：{course.difficulty}　实践占比：{course.practice_ratio}",
        f"对应职业方向：{course.career}",
    ]
    if course.remark:
        lines.append(f"备注：{course.remark}")
    if successors:
        lines.append("培养方案中将本课程列为先修的后续课程：" + "、".join(successors))
    return "\n".join(lines)


def generate_ai_advice(
    major: Major,
    course: Course,
    client,
) -> str:
    """
    生成（或读取缓存的）课程学习建议。

    :return: Markdown 建议文本
    :raises Exception: 透传调用过程中的网络/配置异常，由 UI 层决定降级提示
    """
    key = _advice_cache_key(major.id, course.code)
    cache = load_advice_cache()
    if key in cache and cache[key].get("text"):
        return cache[key]["text"]

    successors = find_successors(course, major.courses)
    messages = [
        {"role": "system", "content": ADVICE_SYSTEM_PROMPT},
        {"role": "user", "content": _build_course_context(course, successors)},
    ]
    chunks = client.chat_stream(messages, temperature=0.6, max_tokens=1500)
    text = "".join(chunks).strip()
    if text:
        cache[key] = {"text": text, "updated_at": time.time()}
        save_advice_cache(cache)
    return text


def fallback_advice(major: Major, course: Course) -> str:
    """无 API Key 或 AI 调用失败时，基于 Excel 字段生成结构化的规则化建议。"""
    successors = find_successors(course, major.courses)
    level = course.difficulty_level
    practice = course.practice_ratio or "—"
    meta = CATEGORY_META.get(course.category, {})

    # —— 重要性 ——
    importance = [
        f"- **体系定位**：本课程属于「{course.category}」，{meta.get('desc', '')}。",
        f"- **职业价值**：主要支撑方向为「{course.career}」，是该方向能力图谱中的关键一环。",
    ]
    if course.remark:
        importance.append(f"- **特别提示**：{course.remark}。")

    # —— 学习路径 ——
    path = []
    if course.prerequisite_list:
        path.append("1. **先修准备**：先掌握 " + "、".join(course.prerequisite_list) + "，再进入本课程。")
    else:
        path.append("1. **先修准备**：本课程无硬性先修要求，可按培养学期直接修读。")
    points = course.content_points[:3]
    focus = ("、".join(points) + " 等") if points else "课程核心知识点"
    path.append(f"2. **课程重点**：重点攻克 {focus}，关注知识体系的系统性。")
    if successors:
        path.append("3. **后续延伸**：本课程是 " + "、".join(successors[:4]) + " 的先修基础，学扎实可为后续课程铺路。")
    else:
        path.append(f"3. **后续延伸**：结合「{course.career}」方向选择实训/项目深化应用。")

    # —— 学习建议 ——
    tips = []
    if level >= 5:
        tips.append("- 课程难度较高（★★★★★），建议课前预习、每周固定复盘，预留充足练习时间。")
    elif level == 4:
        tips.append("- 课程偏难（★★★★），建议跟紧教学进度，遇到疑点及时通过答疑系统解决。")
    else:
        tips.append("- 课程难度适中，保持稳定投入即可，重点放在理解的深度而非死记。")
    try:
        practice_num = int("".join(ch for ch in practice if ch.isdigit()) or "0")
    except ValueError:
        practice_num = 0
    if practice_num >= 40:
        tips.append(f"- 实践占比约 {practice}，务必亲手完成全部实验/项目，能力是在动手过程中形成的。")
    elif practice_num > 0:
        tips.append(f"- 实践占比约 {practice}，在理论学习之外应主动配套上机练习。")
    tips.append(f"- 考核方式为「{course.assessment}」，建议按考核构成及早分配平时与期末的精力。")

    return (
        "### 🎯 重要性分析\n"
        + "\n".join(importance)
        + "\n\n### 🗺️ 学习路径\n"
        + "\n".join(path)
        + "\n\n### 💡 学习建议\n"
        + "\n".join(tips)
        + "\n\n> ⚙️ 以上为基于课程数据表生成的通用建议；配置 DeepSeek API Key 后可获取 AI 个性化建议。"
    )
